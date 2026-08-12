"""Module 9 — Excel Auto-Fill — template storage and fill orchestration."""
import os
import uuid
from pathlib import Path

from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.models.excel_template import ExcelFillLog, ExcelTemplate
from app.utils.excel_autofill import ExcelAutoFillError, fill_template, find_tokens
from app.utils.token_registry import build_token_registry

settings = get_settings()


class TemplateNotFoundError(Exception):
    pass


class ExcelTemplateService:
    def __init__(self, db: AsyncSession):
        self.db = db

    async def upload(self, user_id: uuid.UUID, title: str, original_filename: str, contents: bytes) -> ExcelTemplate:
        tokens = find_tokens(contents)

        os.makedirs(settings.EXCEL_TEMPLATE_DIR, exist_ok=True)
        stored_filename = f"{uuid.uuid4()}.xlsx"
        file_path = Path(settings.EXCEL_TEMPLATE_DIR) / stored_filename
        with open(file_path, "wb") as f:
            f.write(contents)

        from openpyxl import load_workbook
        import io
        wb = load_workbook(io.BytesIO(contents))

        template = ExcelTemplate(
            title=title.strip() or original_filename,
            original_filename=original_filename,
            stored_filename=stored_filename,
            sheet_count=len(wb.sheetnames),
            token_count=len(tokens),
            file_size_bytes=len(contents),
            uploaded_by=user_id,
        )
        self.db.add(template)
        await self.db.commit()
        await self.db.refresh(template, attribute_names=["uploader"])
        return template

    async def delete(self, template_id: uuid.UUID) -> None:
        template = await self.db.get(ExcelTemplate, template_id)
        if not template:
            raise TemplateNotFoundError(f"Template {template_id} not found.")

        file_path = Path(settings.EXCEL_TEMPLATE_DIR) / template.stored_filename
        if file_path.exists():
            try:
                os.remove(file_path)
            except OSError:
                pass

        await self.db.delete(template)
        await self.db.commit()


class ExcelFillService:
    def __init__(self, db: AsyncSession):
        self.db = db

    async def fill(
        self, user_id: uuid.UUID, template_id: uuid.UUID, academic_year: str | None
    ) -> tuple[bytes, str]:
        template = await self.db.get(ExcelTemplate, template_id)
        if not template:
            raise TemplateNotFoundError(f"Template {template_id} not found.")

        file_path = Path(settings.EXCEL_TEMPLATE_DIR) / template.stored_filename
        if not file_path.exists():
            raise TemplateNotFoundError("The template file is missing from storage. Please re-upload it.")

        with open(file_path, "rb") as f:
            contents = f.read()

        tokens = await build_token_registry(self.db, academic_year)

        error_message = None
        filled_bytes = b""
        filled_count = 0
        missing: list[str] = []
        try:
            filled_bytes, filled_count, missing = fill_template(contents, tokens)
        except ExcelAutoFillError as exc:
            error_message = str(exc)

        log = ExcelFillLog(
            template_id=template.id,
            user_id=user_id,
            academic_year=academic_year,
            tokens_filled=filled_count,
            tokens_missing=missing or None,
            error_message=error_message,
            file_size_bytes=len(filled_bytes) if filled_bytes else None,
        )
        self.db.add(log)
        await self.db.commit()

        if error_message:
            raise ExcelAutoFillError(error_message)

        year_slug = (academic_year or "all-years").replace("/", "-")
        filename = f"{Path(template.original_filename).stem}_filled_{year_slug}.xlsx"
        return filled_bytes, filename
